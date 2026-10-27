"""Excel export (right-to-left sheets, Jalali dates)."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .db import DB
from .utils import latin_digits, jdate

KIND_FA = {"expense": "هزینه", "income": "درآمد", "transfer": "انتقال", "adjust": "اصلاح موجودی",
           "debt_in": "دریافت قرض/طلب", "debt_out": "پرداخت قرض/بدهی"}
INTERVAL_FA = {"weekly": "هفتگی", "monthly": "ماهانه", "yearly": "سالانه"}


def _d(value) -> str:
    return latin_digits(jdate(value)) if value else ""


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], money_cols=()):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="264653")
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    for col in money_cols:
        for cell in ws[get_column_letter(col)][1:]:
            cell.number_format = "#,##0"
    for i, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r[i - 1] or "")) for r in rows]) + 4
        ws.column_dimensions[get_column_letter(i)].width = min(width, 45)
    ws.freeze_panes = "A2"


def build_excel(db: DB) -> io.BytesIO:
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "تراکنش‌ها",
           ["شناسه", "تاریخ", "ساعت", "نوع", "مبلغ (تومان)", "حساب", "حساب مقصد", "دسته", "شرح", "منبع"],
           [[t["id"], _d(t["occurred_at"]), t["occurred_at"][11:16], KIND_FA.get(t["kind"], t["kind"]),
             t["amount"], t["account_name"], t["to_account_name"], t["category"], t["description"], t["source"]]
            for t in db.all_transactions()], money_cols=(5,))
    _sheet(wb, "حساب‌ها", ["شناسه", "نام", "بانک", "۴ رقم آخر کارت", "موجودی اولیه", "موجودی فعلی"],
           [[a["id"], a["name"], a["bank"], a["card_last4"], a["initial_balance"], bal] for a, bal in db.balances()],
           money_cols=(5, 6))
    _sheet(wb, "طلب و بدهی", ["شناسه", "شخص", "نوع", "مبلغ", "مانده", "شرح", "تاریخ", "تاریخ تسویه"],
           [[d["id"], d["person"], "طلب" if d["direction"] == "lent" else "بدهی", d["amount"], d["remaining"],
             d["description"], _d(d["created_at"]), _d(d["settled_at"])] for d in db.debts(open_only=False)],
           money_cols=(4, 5))
    _sheet(wb, "اقساط و تکراری", ["شناسه", "عنوان", "مبلغ", "دوره", "سررسید بعدی", "اقساط باقی‌مانده"],
           [[r["id"], r["title"], r["amount"], INTERVAL_FA.get(r["interval"]), _d(r["next_due"]), r["remaining_count"]]
            for r in db.recurring_list()], money_cols=(3,))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
